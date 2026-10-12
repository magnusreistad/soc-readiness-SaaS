"""
Zip-bomb guard for Word uploads: the entry-count and total-size caps must be
checked before python-docx (Strategy 1) or the direct ZIP/XML path
(Strategy 2) decompress anything. Caps are monkeypatched relative to a small
real .docx, so no large files are generated.
"""

import io
import zipfile

import docx
import pytest

from app.utils import document_processor as dp

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT = "Access reviews are performed quarterly"


@pytest.fixture(scope="module")
def docx_bytes() -> bytes:
    # Module-scoped, so it is built before the function-scoped spies below
    # are installed and its own docx.Document() call isn't recorded.
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph(TEXT)
    d.save(buf)
    return buf.getvalue()


def _total_uncompressed(file_bytes: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        return sum(i.file_size for i in zf.infolist())


def _entry_count(file_bytes: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        return len(zf.infolist())


@pytest.fixture
def document_calls(monkeypatch):
    """Records every call to docx.Document while still delegating to it."""
    calls = []
    real = docx.Document

    def spy(*args, **kwargs):
        calls.append(args)
        return real(*args, **kwargs)

    monkeypatch.setattr(docx, "Document", spy)
    return calls


@pytest.fixture
def zip_reads(monkeypatch):
    """Records every member decompressed via ZipFile.read (python-docx and Strategy 2)."""
    reads = []
    real = zipfile.ZipFile.read

    def spy(self, name, *args, **kwargs):
        reads.append(name if isinstance(name, str) else name.filename)
        return real(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", spy)
    return reads


def test_docx_under_caps_extracts_normally(monkeypatch, docx_bytes, document_calls):
    data = docx_bytes
    monkeypatch.setattr(dp, "MAX_ZIP_UNCOMPRESSED_BYTES", _total_uncompressed(data))
    monkeypatch.setattr(dp, "MAX_ZIP_ENTRIES", _entry_count(data))

    result = dp.extract_text(data, "policy.docx", DOCX_MIME)

    assert len(document_calls) == 1
    assert result.file_type == "word"
    assert TEXT in result.text
    assert result.error is None


def test_docx_over_total_size_cap_never_reaches_document(monkeypatch, docx_bytes, document_calls, zip_reads):
    data = docx_bytes
    monkeypatch.setattr(dp, "MAX_ZIP_UNCOMPRESSED_BYTES", _total_uncompressed(data) - 1)

    result = dp.extract_text(data, "policy.docx", DOCX_MIME)

    assert document_calls == []
    assert zip_reads == []          # Strategy 2 skipped too
    assert TEXT not in result.text  # only the non-ZIP fallbacks ran


def test_docx_over_entry_count_cap_never_reaches_document(monkeypatch, docx_bytes, document_calls, zip_reads):
    data = docx_bytes
    monkeypatch.setattr(dp, "MAX_ZIP_ENTRIES", _entry_count(data) - 1)

    result = dp.extract_text(data, "policy.docx", DOCX_MIME)

    assert document_calls == []
    assert zip_reads == []
    assert TEXT not in result.text


def test_direct_zip_fallback_still_works_under_caps(monkeypatch):
    # A .docx missing [Content_Types].xml: python-docx rejects it and the
    # direct ZIP/XML path (Strategy 2) recovers the text.
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", f'<w:document xmlns:w="{w}"><w:t>{TEXT}</w:t></w:document>')
    data = buf.getvalue()
    monkeypatch.setattr(dp, "MAX_ZIP_UNCOMPRESSED_BYTES", _total_uncompressed(data))

    result = dp.extract_text(data, "repacked.docx", DOCX_MIME)

    assert TEXT in result.text
